import html
import io
import json
import urllib.parse
from datetime import datetime
from typing import Any, Dict, List

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from google_auth_oauthlib.flow import InstalledAppFlow, Flow
from googleapiclient.discovery import build
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import importlib
import db
import collector

# Ensure fresh module reloads during development
importlib.reload(db)
importlib.reload(collector)

# Page setup
st.set_page_config(
    page_title="TubeSSOT - YouTube 구독 자산 통합 관리자",
    page_icon="📺",
    layout="wide",
    initial_sidebar_state="expanded",
)

# Initialize DB tables
db.init_db()

# --- SIDEBAR: Settings & Quota Tracking ---
st.sidebar.title("⚙️ 설정 및 API 상태")

# Preload settings from DB
saved_client_id = db.get_setting("google_client_id", "") or ""
saved_client_secret = db.get_setting("google_client_secret", "") or ""

with st.sidebar.expander("🔑 OAuth 클라이언트 자격증명", expanded=not bool(saved_client_id and saved_client_secret)):
    # File uploader for client_secret*.json
    uploaded_secret_file = st.file_uploader(
        "Google Cloud client_secret.json 업로드",
        type=["json"],
        help="Google Cloud Console에서 다운로드한 OAuth 클라이언트 JSON 파일을 업로드하면 ID와 Secret이 자동 입력됩니다.",
    )
    if uploaded_secret_file is not None:
        try:
            content = json.load(uploaded_secret_file)
            installed_data = content.get("installed") or content.get("web") or {}
            c_id = installed_data.get("client_id", "")
            c_secret = installed_data.get("client_secret", "")
            if c_id and c_secret:
                db.set_setting("google_client_id", c_id)
                db.set_setting("google_client_secret", c_secret)
                saved_client_id = c_id
                saved_client_secret = c_secret
                st.success("JSON 파일로부터 자격증명을 성공적으로 불러와 저장했습니다!")
            else:
                st.error("JSON 파일에 유효한 client_id 또는 client_secret이 없습니다.")
        except Exception as e:
            st.error(f"JSON 파일 파싱 실패: {str(e)}")

    client_id_input = st.text_input(
        "Google Client ID",
        value=saved_client_id,
        type="password",
        help="Google Cloud Console -> API 및 서비스 -> 사용자 인증 정보에서 발급받은 OAuth 클라이언트 ID",
    )
    client_secret_input = st.text_input(
        "Google Client Secret",
        value=saved_client_secret,
        type="password",
        help="Google Cloud Console에서 발급받은 OAuth 클라이언트 보안 비밀번호",
    )

    if st.button("자격증명 저장", use_container_width=True):
        db.set_setting("google_client_id", client_id_input.strip())
        db.set_setting("google_client_secret", client_secret_input.strip())
        st.success("자격증명이 데이터베이스에 안전하게 저장되었습니다.")
        st.rerun()

client_id = db.get_setting("google_client_id", "") or ""
client_secret = db.get_setting("google_client_secret", "") or ""

st.sidebar.divider()

# Quota Gauge
st.sidebar.subheader("📊 일일 API 쿼터 사용량")
today_quota = db.get_today_quota()
max_quota = 10000
quota_percent = min(1.0, today_quota / max_quota)

st.sidebar.progress(quota_percent)
col_q1, col_q2 = st.sidebar.columns(2)
col_q1.metric("오늘 소모 쿼터", f"{today_quota:,} U")
col_q2.metric("잔여 쿼터", f"{max(0, max_quota - today_quota):,} U")

if st.sidebar.button("🔄 쿼터/화면 새로고침", use_container_width=True):
    st.rerun()

recent_logs = db.get_recent_quota_logs(limit=5)
if recent_logs:
    with st.sidebar.expander("📋 최근 API 소모 내역 (5건)"):
        for log in recent_logs:
            st.caption(f"• `{log['created_at'][11:19]}`: +**{log['units']}U** ({log['action']})")

st.sidebar.caption(
    f"기본 무료 할당량(10,000 유닛)의 **{quota_percent * 100:.1f}%** 소모됨.\n"
    "TubeSSOT는 재생목록 캐시(`UU...`) 및 50개 배치 조회를 활용해 극단적인 쿼터 효율성을 보장합니다."
)

st.sidebar.divider()

# Connected channels in sidebar
tokens = db.get_all_tokens()
st.sidebar.markdown("### 🔗 연동된 YouTube 채널")
if tokens:
    st.sidebar.caption(f"총 **{len(tokens)}개** 채널 연동됨")
    for t in tokens:
        st.sidebar.success(f"📺 **{t['channel_title']}**\n\n`{t['channel_id']}`")
else:
    st.sidebar.warning("⚠️ 연동된 채널이 없습니다.\n\n'1. 연동 계정 관리' 탭에서 먼저 계정을 연동해 주세요.")

st.sidebar.divider()
st.sidebar.info("💡 **TubeSSOT v1.0**\n단일 진실 공급원(SSOT) 기반 YouTube 구독 자산 및 수명주기 관리자")


# --- OAuth Callback Handler (Redirect from Google) ---
if "error" in st.query_params:
    oauth_err = st.query_params.get("error")
    st.session_state["oauth_status"] = ("error", f"Google 인증이 취소되었거나 오류가 발생했습니다: {oauth_err}")
    st.query_params.clear()
    st.rerun()

if "code" in st.query_params and client_id and client_secret:
    auth_code = st.query_params["code"]
    print(f"[OAuth Callback] Received auth code, exchanging for tokens...", flush=True)
    with st.spinner("Google 계정 인증 정보를 등록하는 중입니다..."):
        try:
            client_config = {
                "installed": {
                    "client_id": client_id,
                    "client_secret": client_secret,
                    "auth_uri": "https://accounts.google.com/o/oauth2/auth",
                    "token_uri": "https://oauth2.googleapis.com/token",
                    "redirect_uris": ["http://localhost:8501/", "http://localhost:8501", "http://localhost:8080/"],
                }
            }
            # Attempt exchange with http://localhost:8501/
            flow = None
            for r_uri in ["http://localhost:8501/", "http://localhost:8501"]:
                try:
                    flow = Flow.from_client_config(
                        client_config,
                        scopes=collector.SCOPES,
                        redirect_uri=r_uri,
                    )
                    flow.autogenerate_code_verifier = False
                    flow.fetch_token(code=auth_code)
                    break
                except Exception as ex_token:
                    print(f"[OAuth Callback] Fetch token with {r_uri} failed: {ex_token}", flush=True)

            if not flow or not flow.credentials:
                raise RuntimeError("Google 서버로부터 토큰을 획득하지 못했습니다.")

            creds = flow.credentials
            yt = build("youtube", "v3", credentials=creds, cache_discovery=False)
            my_channels = yt.channels().list(part="snippet", mine=True).execute()

            if my_channels.get("items"):
                ch_info = my_channels["items"][0]
                c_id = ch_info["id"]
                c_title = ch_info["snippet"]["title"]
            else:
                # Fallback: Google Account without dedicated YouTube channel
                c_id = f"user_{creds.client_id[:12]}"
                c_title = "Google Linked Account (기본 계정)"

            db.upsert_token(
                channel_id=c_id,
                email="Linked Google Account",
                title=c_title,
                refresh_token=creds.refresh_token,
                access_token=creds.token,
                expiry=creds.expiry,
            )
            print(f"[OAuth Callback] Successfully linked channel: {c_title} ({c_id})", flush=True)
            st.session_state["oauth_status"] = ("success", f"🎉 채널 **'{c_title}'** ({c_id}) 연동이 성공적으로 등록되었습니다!")
        except Exception as e:
            err_msg = f"계정 연동 실패: {str(e)}"
            print(f"[OAuth Callback Error] {err_msg}", flush=True)
            st.session_state["oauth_status"] = ("error", err_msg)
        finally:
            st.query_params.clear()
            st.rerun()


# Display persistent OAuth Status Message if present
if "oauth_status" in st.session_state:
    status_type, status_text = st.session_state.pop("oauth_status")
    if status_type == "success":
        st.success(status_text)
        st.balloons()
    else:
        st.error(status_text)

# Refresh tokens list
tokens = db.get_all_tokens()

# Main Top Status Banner
if tokens:
    connected_names = " | ".join([f"**{t['channel_title']}** (`{t['channel_id']}`)" for t in tokens])
    st.info(f"🔗 **현재 연동된 YouTube 채널 ({len(tokens)}개):** {connected_names}")
else:
    st.warning("⚠️ **연동된 YouTube 채널이 없습니다.** 아래 '1. 👥 연동 계정 관리' 탭에서 계정을 연동해 주세요.")


# --- HELPER FUNCTIONS & UI COMPONENTS ---
def format_count(val: Any, unit: str = "") -> str:
    if val is None or val == "":
        return "-"
    try:
        num = int(val)
        if num >= 100_000_000:
            res = f"{num / 100_000_000:.1f}억"
        elif num >= 10_000:
            res = f"{num / 10_000:.1f}만"
        elif num >= 1_000:
            res = f"{num / 1_000:.1f}천"
        else:
            res = f"{num:,}"
        return f"{res}{unit}" if unit else res
    except (ValueError, TypeError):
        return f"{val}{unit}" if unit else str(val)


status_color_map = {
    "GREEN": ("🟢 활성", "#0F9D58"),
    "YELLOW": ("🟡 정체", "#E37400"),
    "RED": ("🔴 휴면", "#D93025"),
    "UNKNOWN": ("⚪ 미확인", "#70757A"),
}

status_badge_map = {
    "GREEN": "🟢 활성",
    "YELLOW": "🟡 정체",
    "RED": "🔴 휴면",
    "UNKNOWN": "⚪ 미확인",
}


def render_channel_card_header(ch: dict, view_type: str = "inbox") -> str:
    """Renders a pixel-perfect 76px header matching thumbnail height with status above 1-line title."""
    escaped_title = html.escape(ch.get("title") or "제목 없음")
    raw_handle = ch.get("custom_url") or (f"ID: {ch['channel_id'][:14]}..." if ch.get("channel_id") else "-")
    escaped_handle = html.escape(raw_handle)
    thumb_url = ch.get("thumbnail_url") or ""

    status = ch.get("liveness_status", "UNKNOWN")
    badge_label, badge_color = status_color_map.get(status, ("⚪ 미확인", "#70757A"))

    if view_type == "keep":
        badge_html = f"<span style='color: #0F9D58; font-weight: 700; font-size: 0.78rem;'>💚 유지</span> <span style='font-size: 0.72rem; color: {badge_color};'>({badge_label})</span>"
    elif view_type == "archive":
        badge_html = f"<span style='color: #888888; font-weight: 700; font-size: 0.78rem;'>🗄️ 보류</span> <span style='font-size: 0.72rem; color: {badge_color};'>({badge_label})</span>"
    else:
        badge_html = f"<span style='color: {badge_color}; font-weight: 700; font-size: 0.78rem;'>{badge_label}</span>"

    thumb_html = (
        f"<img src='{html.escape(thumb_url)}' style='width: 76px; height: 76px; object-fit: cover;' onerror=\"this.style.display='none';\" />"
        if thumb_url
        else "<div style='font-size: 2rem;'>📺</div>"
    )

    return f"""
    <div style="display: flex; align-items: center; gap: 10px; height: 76px; margin-bottom: 2px;">
        <div style="width: 76px; height: 76px; min-width: 76px; border-radius: 8px; overflow: hidden; background: #2b2b2b; display: flex; align-items: center; justify-content: center; flex-shrink: 0;">
            {thumb_html}
        </div>
        <div style="display: flex; flex-direction: column; justify-content: space-between; height: 76px; min-width: 0; flex: 1; overflow: hidden; padding: 2px 0;">
            <div style="line-height: 1.2;">
                {badge_html}
            </div>
            <div style="font-size: 0.95rem; font-weight: 700; line-height: 1.3; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; color: inherit;" title="{escaped_title}">
                {escaped_title}
            </div>
            <div style="font-size: 0.74rem; color: #888888; line-height: 1.2; white-space: nowrap; overflow: hidden; text-overflow: ellipsis;" title="{escaped_handle}">
                {escaped_handle}
            </div>
        </div>
    </div>
    """


@st.dialog("📺 채널 상세 정보 및 최근 영상 (최대 5개)", width="large")
def show_channel_modal(ch: dict):
    # 상단: 채널 정보
    c1, c2 = st.columns([1, 4])
    with c1:
        if ch.get("thumbnail_url"):
            st.image(ch["thumbnail_url"], width=120)
        else:
            st.markdown("📺")
    with c2:
        badge = status_badge_map.get(ch.get("liveness_status"), "")
        rev_status = ch.get("review_status", "INBOX")
        status_tag = ""
        if rev_status == "KEEP":
            status_tag = " [💚 유지 상태]"
        elif rev_status == "ARCHIVED":
            status_tag = " [🗄️ 보류 상태]"
        else:
            status_tag = " [📥 미분류 상태]"
        st.markdown(f"### **{ch.get('title')}** {badge}{status_tag}")
        direct_url = f"https://www.youtube.com/{ch['custom_url']}" if ch.get("custom_url") else f"https://www.youtube.com/channel/{ch.get('channel_id')}"
        sub_url = f"{direct_url}?sub_confirmation=1"
        st.markdown(f"[🌐 YouTube 채널 바로가기]({direct_url}) ｜ [🔔 원클릭 구독/확인 링크]({sub_url})")

    # 4 Key Metrics Bar
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("구독자 수", ch.get("subscriber_count_str", "-"))
    m2.metric("전체 영상 수", ch.get("video_count_str", "-"))
    m3.metric("최근 업로드일", ch.get("last_upload_at") or "-")
    days = ch.get("days_since_last_upload")
    m4.metric("미업로드 경과일수", f"{days:,}일 전" if days is not None else "-")

    if ch.get("categories"):
        st.markdown(f"🏷️ **카테고리:** {', '.join(ch['categories'])}")
    if ch.get("source_accounts"):
        st.markdown(f"👤 **구독 소속 계정:** {', '.join(ch['source_accounts'])}")
    if ch.get("description"):
        with st.expander("📝 채널 소개글 보기", expanded=False):
            st.write(ch["description"])

    # Quick status change in modal
    st.markdown("##### ⚡ 분류 상태 변경:")
    sm1, sm2, sm3 = st.columns(3)
    with sm1:
        if st.button("💚 [유지]로 지정", key=f"mdl_keep_{ch['channel_id']}", disabled=(rev_status == "KEEP"), use_container_width=True):
            db.set_channel_keep(ch["channel_id"])
            st.toast(f"'{ch['title']}' 채널이 유지 목록으로 지정되었습니다.", icon="💚")
            st.rerun()
    with sm2:
        if st.button("📦 [보류]로 이동", key=f"mdl_arch_{ch['channel_id']}", disabled=(rev_status == "ARCHIVED"), use_container_width=True):
            db.set_channel_archived(ch["channel_id"], True)
            st.toast(f"'{ch['title']}' 채널이 보류 보관함으로 이동되었습니다.", icon="📦")
            st.rerun()
    with sm3:
        if st.button("↩️ [미분류]로 복귀", key=f"mdl_inbox_{ch['channel_id']}", disabled=(rev_status == "INBOX"), use_container_width=True):
            db.set_channel_status(ch["channel_id"], "INBOX")
            st.toast(f"'{ch['title']}' 채널이 미분류 상태로 복귀되었습니다.", icon="↩️")
            st.rerun()

    st.divider()

    # 하단: 5개의 최근 동영상 목록
    st.markdown("#### 🎬 최근 업로드 동영상 (최대 5개)")
    recent_vids = ch.get("recent_videos", [])
    if recent_vids:
        num_cols = min(5, len(recent_vids))
        v_cols = st.columns(num_cols)
        for v_idx, v in enumerate(recent_vids[:num_cols]):
            with v_cols[v_idx]:
                with st.container(border=True):
                    if v.get("thumbnail_url"):
                        st.image(v.get("thumbnail_url"), use_container_width=True)
                    v_title = v.get("title", "제목 없음")
                    v_url = v.get("video_url", f"https://www.youtube.com/watch?v={v.get('video_id')}")
                    st.markdown(f"**[{v_title}]({v_url})**")
                    pub = v.get("published_at", "")[:10]
                    if pub:
                        st.caption(f"📅 {pub}")
    else:
        st.info("이 채널에 수집된 최근 영상이 없습니다.")

    with st.expander("🛠️ API 원본 JSON 페이로드 (Raw JSON) 검사"):
        st.json({
            "raw_subscription": ch.get("raw_subscription_json"),
            "raw_channel": ch.get("raw_channel_json"),
        })


def render_channel_cards_grid(
    items: List[Dict[str, Any]],
    view_type: str = "inbox",  # "inbox", "keep", "archive"
    page_key_prefix: str = "inbox",
):
    """Renders a responsive 3-column card grid with infinite scroll on user scroll."""
    if not items:
        if view_type == "inbox":
            st.success("🎉 **미분류 채널이 없습니다 (Inbox Zero 달성)!**\n\n모든 구독 채널이 '유지' 또는 '보류'로 성공적으로 분류되었습니다. 새로운 계정/채널을 동기화하면 새 구독이 이곳에 나타납니다.")
        elif view_type == "keep":
            st.info("ℹ️ 현재 유지 목록으로 분류된 채널이 없습니다. 상단 **'3. 📥 미분류 채널'** 화면에서 유지할 채널의 **[💚 유지]** 버튼을 눌러보세요.")
        elif view_type == "archive":
            st.info("ℹ️ 현재 보류된 채널이 없습니다. 불필요하거나 정리가 필요한 채널 카드의 **[📦 보류]** 버튼을 누르면 이 보관함으로 이동합니다.")
        return

    batch_size = 24
    total_cards = len(items)

    visible_key = f"{page_key_prefix}_visible_count"
    if visible_key not in st.session_state:
        st.session_state[visible_key] = batch_size

    # Limit slice to available items
    current_visible = min(total_cards, st.session_state[visible_key])
    page_items = items[:current_visible]

    # Live progress caption
    st.caption(f"📊 현재 **{current_visible:,}개** / 전체 **{total_cards:,}개** 채널 노출 중 (아래로 스크롤 시 자동 추가 로딩)")

    cards_per_row = 3
    for r_idx in range(0, len(page_items), cards_per_row):
        row_cols = st.columns(cards_per_row)
        for c_idx in range(cards_per_row):
            item_idx = r_idx + c_idx
            if item_idx < len(page_items):
                ch = page_items[item_idx]
                with row_cols[c_idx]:
                    with st.container(border=True):
                        # 1. Uniform Header (Status on top + 1-line title + handle, matching 76px thumbnail)
                        st.markdown(render_channel_card_header(ch, view_type=view_type), unsafe_allow_html=True)

                        # 2. Statistics (Subscribers & Total Videos)
                        st.divider()
                        st_col1, st_col2 = st.columns(2)
                        st_col1.markdown(f"👥 **구독자:** {ch.get('subscriber_count_str', '-')}")
                        st_col2.markdown(f"🎬 **영상:** {ch.get('video_count_str', '-')}")

                        # 3. Liveness / upload activity
                        days_txt = f"{ch['days_since_last_upload']}일 전" if ch.get("days_since_last_upload") is not None else "-"
                        st.caption(f"📅 최근 업로드: {ch.get('last_upload_at', '-')[:10]} ({days_txt})")

                        # Direct YouTube URL
                        direct_url = f"https://www.youtube.com/{ch['custom_url']}" if ch.get("custom_url") else f"https://www.youtube.com/channel/{ch['channel_id']}"

                        # 4. Action Row 1: View / Modal & YouTube Link
                        act1_col1, act1_col2 = st.columns([1.5, 1])
                        with act1_col1:
                            if st.button("🔍 상세 / 5영상", key=f"btn_modal_{page_key_prefix}_{ch['channel_id']}", type="primary", use_container_width=True):
                                show_channel_modal(ch)
                        with act1_col2:
                            st.link_button("🌐 바로가기", direct_url, use_container_width=True)

                        # 5. Action Row 2: Status actions (3 equal columns)
                        act2_c1, act2_c2, act2_c3 = st.columns(3)
                        if view_type == "inbox":
                            with act2_c1:
                                if st.button("💚 유지", key=f"btn_keep_{page_key_prefix}_{ch['channel_id']}", use_container_width=True, help="채널을 유지 목록으로 분류합니다"):
                                    db.set_channel_keep(ch["channel_id"])
                                    st.toast(f"'{ch['title']}' 채널이 유지 목록으로 이동되었습니다.", icon="💚")
                                    st.rerun()
                            with act2_c2:
                                if st.button("📦 보류", key=f"btn_arch_{page_key_prefix}_{ch['channel_id']}", use_container_width=True, help="채널을 보류 보관함으로 이동합니다"):
                                    db.set_channel_archived(ch["channel_id"], True)
                                    st.toast(f"'{ch['title']}' 채널이 보류 보관함으로 이동되었습니다.", icon="📦")
                                    st.rerun()
                            with act2_c3:
                                with st.popover("❌ 해지", use_container_width=True):
                                    st.markdown(f"**'{ch['title']}'** 구독 해지")
                                    st.caption("YouTube 채널 페이지에서 구독 취소를 확정합니다.")
                                    st.link_button("👉 YouTube 채널 열기", direct_url, use_container_width=True)
                                    st.divider()
                                    st.caption("TubeSSOT 로컬 DB 목록에서 이 채널을 영구 제외합니다.")
                                    if st.button("🗑️ DB에서 영구 삭제", key=f"del_ch_{page_key_prefix}_{ch['channel_id']}", type="secondary", use_container_width=True):
                                        db.delete_master_channel(ch["channel_id"])
                                        st.toast(f"'{ch['title']}' 채널이 DB에서 삭제되었습니다.", icon="🗑️")
                                        st.rerun()

                        elif view_type == "keep":
                            with act2_c1:
                                if st.button("↩️ 미분류", key=f"btn_inbox_{page_key_prefix}_{ch['channel_id']}", use_container_width=True, help="채널을 다시 미분류(검토 대기) 상태로 복귀합니다"):
                                    db.set_channel_status(ch["channel_id"], "INBOX")
                                    st.toast(f"'{ch['title']}' 채널이 미분류 상태로 복귀되었습니다.", icon="↩️")
                                    st.rerun()
                            with act2_c2:
                                if st.button("📦 보류", key=f"btn_arch_{page_key_prefix}_{ch['channel_id']}", use_container_width=True, help="채널을 보류 보관함으로 이동합니다"):
                                    db.set_channel_archived(ch["channel_id"], True)
                                    st.toast(f"'{ch['title']}' 채널이 보류 보관함으로 이동되었습니다.", icon="📦")
                                    st.rerun()
                            with act2_c3:
                                with st.popover("❌ 해지", use_container_width=True):
                                    st.markdown(f"**'{ch['title']}'** 구독 해지")
                                    st.caption("YouTube 채널 페이지에서 구독 취소를 확정합니다.")
                                    st.link_button("👉 YouTube 채널 열기", direct_url, use_container_width=True)
                                    st.divider()
                                    st.caption("TubeSSOT 로컬 DB 목록에서 이 채널을 영구 제외합니다.")
                                    if st.button("🗑️ DB에서 영구 삭제", key=f"del_ch_{page_key_prefix}_{ch['channel_id']}", type="secondary", use_container_width=True):
                                        db.delete_master_channel(ch["channel_id"])
                                        st.toast(f"'{ch['title']}' 채널이 DB에서 삭제되었습니다.", icon="🗑️")
                                        st.rerun()

                        else:  # view_type == "archive"
                            with act2_c1:
                                if st.button("💚 유지", key=f"btn_keep_{page_key_prefix}_{ch['channel_id']}", use_container_width=True, help="채널을 유지 목록으로 이동합니다"):
                                    db.set_channel_keep(ch["channel_id"])
                                    st.toast(f"'{ch['title']}' 채널이 유지 목록으로 이동되었습니다.", icon="💚")
                                    st.rerun()
                            with act2_c2:
                                if st.button("↩️ 미분류", key=f"btn_restore_{page_key_prefix}_{ch['channel_id']}", use_container_width=True, help="채널을 미분류 검토 목록으로 복원합니다"):
                                    db.set_channel_status(ch["channel_id"], "INBOX")
                                    st.toast(f"'{ch['title']}' 채널이 미분류 상태로 복원되었습니다!", icon="♻️")
                                    st.rerun()
                            with act2_c3:
                                with st.popover("❌ 영구 삭제", use_container_width=True):
                                    st.markdown(f"**'{ch['title']}'** 완전 삭제")
                                    st.caption("YouTube 채널 페이지에서 구독 취소 여부를 확인합니다.")
                                    st.link_button("👉 YouTube 채널 열기", direct_url, use_container_width=True)
                                    st.divider()
                                    st.caption("TubeSSOT 로컬 DB에서 이 채널을 완전히 삭제합니다.")
                                    if st.button("🗑️ DB에서 영구 삭제", key=f"del_arch_{page_key_prefix}_{ch['channel_id']}", type="secondary", use_container_width=True):
                                        db.delete_master_channel(ch["channel_id"])
                                        st.toast(f"'{ch['title']}' 채널이 DB에서 삭제되었습니다.", icon="🗑️")
                                        st.rerun()

    # Infinite Scroll Bottom Area
    if current_visible < total_cards:
        remaining = total_cards - current_visible
        next_count = min(batch_size, remaining)
        sentinel_id = f"sentinel_{page_key_prefix}"

        # Sentinel element for IntersectionObserver
        st.markdown(
            f"""
            <div id="{sentinel_id}" style="height: 40px; margin: 24px 0 10px 0; text-align: center; display: flex; align-items: center; justify-content: center; background: rgba(255,255,255,0.03); border-radius: 8px; border: 1px dashed rgba(255,255,255,0.18);">
                <span style="color: #aaa; font-size: 0.88rem;">⏳ 아래로 스크롤하면 다음 {next_count}개 채널을 자동으로 불러옵니다...</span>
            </div>
            """,
            unsafe_allow_html=True,
        )

        # Fallback load more button
        load_col1, load_col2, load_col3 = st.columns([1, 2, 1])
        with load_col2:
            if st.button(
                f"⬇️ {next_count}개 더 불러오기 ({current_visible:,} / {total_cards:,})",
                key=f"btn_more_{page_key_prefix}",
                use_container_width=True,
            ):
                st.session_state[visible_key] = current_visible + batch_size
                st.rerun()

        # JavaScript IntersectionObserver for butter-smooth automatic infinite scrolling
        components.html(
            f"""
            <script>
            (function() {{
                try {{
                    const parentWin = window.parent;
                    const parentDoc = parentWin.document;
                    const sentinel = parentDoc.getElementById("{sentinel_id}");

                    parentWin.__isLoadingMore = false;

                    if (sentinel && parentWin.IntersectionObserver) {{
                        const observer = new parentWin.IntersectionObserver((entries) => {{
                            entries.forEach(entry => {{
                                if (entry.isIntersecting && !parentWin.__isLoadingMore) {{
                                    const buttons = Array.from(parentDoc.querySelectorAll('button'));
                                    const loadBtn = buttons.find(b => 
                                        b.offsetParent !== null && 
                                        (b.innerText || b.textContent || '').includes('더 불러오기')
                                    );
                                    if (loadBtn) {{
                                        parentWin.__isLoadingMore = true;
                                        observer.disconnect();
                                        loadBtn.click();
                                    }}
                                }}
                            }});
                        }}, {{
                            root: null,
                            rootMargin: '400px',
                            threshold: 0.01
                        }});

                        observer.observe(sentinel);
                    }}
                }} catch(err) {{
                    console.error("Infinite scroll error:", err);
                }}
            }})();
            </script>
            """,
            height=0,
            width=0,
        )
    else:
        st.markdown(
            f"""
            <div style="text-align: center; padding: 30px 0 10px 0; color: #888;">
                <div style="font-size: 1.4rem; margin-bottom: 4px;">🎉</div>
                <div style="font-weight: 600; font-size: 0.95rem;">모든 채널 ({total_cards:,}개)을 다 불러왔습니다.</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        top_c1, top_c2, top_c3 = st.columns([1.5, 1, 1.5])
        with top_c2:
            if st.button("⬆️ 맨 위로 이동", key=f"btn_top_{page_key_prefix}", use_container_width=True):
                components.html(
                    """
                    <script>
                    try {
                        window.parent.scrollTo({top: 0, behavior: 'smooth'});
                    } catch(e) {}
                    </script>
                    """,
                    height=0,
                    width=0,
                )


# --- PRELOAD MASTER RECORDS FOR DASHBOARD, TABS & EXPORT ---
records = db.fetch_master_records()
data_list = []
all_categories = set()
all_sources = set()

for r in records:
    cats = json.loads(r["categories"]) if r["categories"] else []
    sources = json.loads(r["source_accounts"]) if r["source_accounts"] else []
    raw_ch = json.loads(r["raw_channel_json"]) if r["raw_channel_json"] else {}
    stats = raw_ch.get("statistics", {})
    sub_count = stats.get("subscriberCount")
    video_count = stats.get("videoCount")
    view_count = stats.get("viewCount")
    is_arch = bool(r["is_archived"]) if ("is_archived" in r.keys() and r["is_archived"]) else False
    rev_status = r["review_status"] if ("review_status" in r.keys() and r["review_status"]) else ("ARCHIVED" if is_arch else "INBOX")

    for c in cats:
        all_categories.add(c)
    for s in sources:
        all_sources.add(s)

    data_list.append({
        "channel_id": r["channel_id"],
        "liveness_status": r["liveness_status"] or "UNKNOWN",
        "review_status": rev_status,
        "title": r["title"],
        "categories": cats,
        "categories_str": ", ".join(cats),
        "days_since_last_upload": r["days_since_last_upload"],
        "last_upload_at": r["last_upload_at"] or "-",
        "source_accounts": sources,
        "source_accounts_str": ", ".join(sources),
        "custom_url": r["custom_url"] or "",
        "description": r["description"] or "",
        "thumbnail_url": r["thumbnail_url"] or "",
        "subscriber_count": sub_count,
        "video_count": video_count,
        "view_count": view_count,
        "subscriber_count_str": format_count(sub_count, "명"),
        "video_count_str": format_count(video_count, "개"),
        "is_archived": is_arch,
        "recent_videos": json.loads(r["recent_videos"]) if r["recent_videos"] else [],
        "raw_subscription_json": json.loads(r["raw_subscription_json"]) if r["raw_subscription_json"] else {},
        "raw_channel_json": raw_ch,
    })

master_df = pd.DataFrame(data_list)
inbox_df = master_df[master_df["review_status"] == "INBOX"] if not master_df.empty else pd.DataFrame()
keep_df = master_df[master_df["review_status"] == "KEEP"] if not master_df.empty else pd.DataFrame()
archived_df = master_df[master_df["review_status"] == "ARCHIVED"] if not master_df.empty else pd.DataFrame()


# --- MAIN CONTENT TABS (Static labels ensure active tab is preserved across reruns) ---
tab_accounts, tab_sync, tab_inbox, tab_keep, tab_archive, tab_export = st.tabs([
    "1. 👥 연동 계정",
    "2. 🔄 수집 및 동기화",
    "3. 📥 미분류 채널 (Inbox)",
    "4. 💚 유지 채널 (Keep)",
    "5. 🗄️ 보류 보관함 (Archive)",
    "6. 📥 엑셀 내보내기",
])

# ==========================================
# TAB 1: 연동 계정 관리 (OAuth Token CRUD)
# ==========================================
with tab_accounts:
    st.subheader("연동된 YouTube 채널 목록")
    st.caption("개인 계정, 업무용 계정, 스터디용 브랜드 채널 등 여러 계정을 등록하여 통합 관리할 수 있습니다.")

    if tokens:
        for idx, t in enumerate(tokens):
            with st.container(border=True):
                col_info, col_id, col_btn = st.columns([4, 3, 1])
                with col_info:
                    st.markdown(f"### 📺 **{t['channel_title']}**")
                    st.caption(f"이메일: `{t['account_email']}` | 등록일: {t['created_at']}")
                with col_id:
                    st.code(t["channel_id"], language="text")
                with col_btn:
                    if st.button("연결 해제", key=f"del_{t['channel_id']}", type="secondary", use_container_width=True):
                        db.delete_token(t["channel_id"])
                        st.toast(f"'{t['channel_title']}' 채널 연동이 해제되었습니다.")
                        st.rerun()
    else:
        st.info("현재 등록된 YouTube 계정/채널이 없습니다. 아래에서 신규 채널을 연동해 주세요.")

    st.divider()
    st.subheader("신규 YouTube 채널 연동 (OAuth 2.0)")

    if not client_id or not client_secret:
        st.warning("⚠️ 사이드바에서 먼저 **Google Client ID**와 **Google Client Secret**을 입력하고 저장해 주세요.")
    else:
        client_config = {
            "installed": {
                "client_id": client_id,
                "client_secret": client_secret,
                "auth_uri": "https://accounts.google.com/o/oauth2/auth",
                "token_uri": "https://oauth2.googleapis.com/token",
                "redirect_uris": ["http://localhost:8501/", "http://localhost:8501", "http://localhost:8080/"],
            }
        }

        try:
            flow = Flow.from_client_config(
                client_config,
                scopes=collector.SCOPES,
                redirect_uri="http://localhost:8501/",
            )
            flow.autogenerate_code_verifier = False
            auth_url, _ = flow.authorization_url(prompt="consent", access_type="offline")

            st.markdown("""
            **원클릭 계정 연동 안내:**  
            아래 버튼을 클릭하면 Google 로그인 및 채널 선택 창이 열립니다.  
            구독 목록을 가져올 **계정 또는 브랜드 채널**을 선택하고 권한을 허용하시면, 자동으로 이 대시보드(`localhost:8501`)로 돌아오며 즉시 연동됩니다.
            """)

            st.link_button(
                "🚀 [원클릭] Google 계정 로그인 및 채널 선택하기",
                auth_url,
                type="primary",
                use_container_width=True,
            )
            st.caption("ℹ️ 클릭 시 Google 로그인 화면으로 이동합니다.")

        except Exception as e:
            st.error(f"인증 URL 생성 실패: {str(e)}")

        st.markdown("---")
        with st.expander("🛠️ 수동 인증 또는 다른 포트(8080) 리다이렉트 붙여넣기"):
            st.markdown("""
            Google 로그인 완료 후 브라우저 주소창이 `http://localhost:8080/?code=...` 등으로 이동하여 페이지가 열리지 않거나,
            직접 인증 코드를 복사하신 경우 아래에 붙여넣어 수동 등록하실 수 있습니다.
            """)
            col_m1, col_m2 = st.columns([3, 1])
            with col_m1:
                manual_input = st.text_input(
                    "주소창 전체 URL 또는 코드 붙여넣기:",
                    placeholder="http://localhost:8080/?code=4/0A... 또는 4/0A...",
                )
            with col_m2:
                redirect_choice = st.selectbox(
                    "사용한 리다이렉트 URI:",
                    ["http://localhost:8501/", "http://localhost:8501", "http://localhost:8080/"],
                )

            if st.button("수동 연동 등록 완료", use_container_width=True):
                if manual_input.strip():
                    try:
                        raw_val = manual_input.strip()
                        code = raw_val
                        if "code=" in raw_val:
                            parsed_qs = urllib.parse.parse_qs(urllib.parse.urlparse(raw_val).query)
                            code = parsed_qs.get("code", [raw_val])[0]

                        manual_flow = Flow.from_client_config(
                            client_config,
                            scopes=collector.SCOPES,
                            redirect_uri=redirect_choice,
                        )
                        manual_flow.autogenerate_code_verifier = False
                        manual_flow.fetch_token(code=code)
                        creds = manual_flow.credentials
                        yt = build("youtube", "v3", credentials=creds, cache_discovery=False)
                        my_channels = yt.channels().list(part="snippet", mine=True).execute()

                        if my_channels.get("items"):
                            ch_info = my_channels["items"][0]
                            c_id = ch_info["id"]
                            c_title = ch_info["snippet"]["title"]
                        else:
                            c_id = f"user_{creds.client_id[:12]}"
                            c_title = "Google Linked Account (기본 계정)"

                        db.upsert_token(
                            channel_id=c_id,
                            email="Linked Google Account",
                            title=c_title,
                            refresh_token=creds.refresh_token,
                            access_token=creds.token,
                            expiry=creds.expiry,
                        )
                        st.session_state["oauth_status"] = ("success", f"🎉 채널 **'{c_title}'** ({c_id}) 수동 연동이 성공적으로 등록되었습니다!")
                        st.rerun()
                    except Exception as e:
                        st.error(f"수동 연동 실패: {str(e)}")
                else:
                    st.warning("URL 또는 코드를 입력해 주세요.")


# ==========================================
# TAB 2: 데이터 수집 및 동기화 (Sync Runner)
# ==========================================
with tab_sync:
    st.subheader("데이터 동기화 오케스트레이터")
    st.write(
        "등록된 모든 연동 계정의 구독 목록을 순회 수집하여 `channel_id` 기준으로 중복을 병합하고,\n"
        "각 채널의 최신 영상 3개를 검사하여 활동성(Liveness: 🟢활성 / 🟡정체 / 🔴휴면 / ⚪미확인)을 판정합니다."
    )

    tokens = db.get_all_tokens()
    if not tokens:
        st.warning("⚠️ **등록된 연동 계정이 없습니다.** '1. 👥 연동 계정 관리' 탭으로 이동하여 YouTube 채널을 먼저 연동해 주세요.")
    elif not client_id or not client_secret:
        st.warning("⚠️ 사이드바에서 Client ID와 Secret을 먼저 설정하세요.")
    else:
        st.markdown("#### 📋 동기화 대상 브랜드 채널 선택")
        st.caption("여러 계정이나 브랜드 채널을 등록했을 때, 특정 채널의 구독 목록만 선택적으로 수집하여 시간과 API 쿼터 소모를 방지할 수 있습니다.")

        channel_map = {t["channel_id"]: t["channel_title"] for t in tokens}
        all_channel_ids = list(channel_map.keys())

        # Quick select buttons
        q_col1, q_col2, q_col3 = st.columns([1, 1, 3])
        with q_col1:
            if st.button("모두 선택", key="btn_sync_all_channels", use_container_width=True):
                st.session_state["sync_target_channels_selector"] = all_channel_ids
                st.rerun()
        with q_col2:
            if st.button("선택 해제", key="btn_sync_clear_channels", use_container_width=True):
                st.session_state["sync_target_channels_selector"] = []
                st.rerun()

        selected_channels = st.multiselect(
            "구독 목록을 수집할 브랜드 채널을 선택하세요:",
            options=all_channel_ids,
            default=all_channel_ids,
            format_func=lambda cid: f"📺 {channel_map.get(cid, cid)} ({cid[:12]}...)",
            key="sync_target_channels_selector",
            help="선택한 채널들의 구독 목록만 순회 수집하여 기존 DB와 통합합니다.",
        )

        if not selected_channels:
            st.warning("⚠️ 최소 1개 이상의 브랜드 채널을 선택해야 동기화를 실행할 수 있습니다.")
        else:
            selected_names = [f"**{channel_map[cid]}**" for cid in selected_channels if cid in channel_map]
            st.info(f"🎯 **선택된 동기화 대상 ({len(selected_channels)}개):** {' ｜ '.join(selected_names)}")

        st.markdown("---")
        st.markdown("#### ⚙️ 동기화 범위 선택 (개발/운영 모드)")
        col_m1, col_m2 = st.columns([3, 2])
        with col_m1:
            range_option = st.selectbox(
                "분석 대상 채널 수 제한:",
                [
                    "전체 채널 수집 (All Channels - 운영 모드)",
                    "빠른 테스트 (10개만 - 초고속 개발용)",
                    "샘플 수집 (50개)",
                    "중간 수집 (100개)",
                    "직접 개수 입력",
                ],
                index=0,
                help="개발 중에는 10개 또는 50개만 부분 수집하여 시간과 API 쿼터를 대폭 절약할 수 있습니다.",
            )
        with col_m2:
            if range_option == "빠른 테스트 (10개만 - 초고속 개발용)":
                max_ch = 10
                st.info("⚡ 약 10~15 쿼터 소모 (약 1~2초 소요)")
            elif range_option == "샘플 수집 (50개)":
                max_ch = 50
                st.info("⚡ 약 55 쿼터 소모 (약 3~5초 소요)")
            elif range_option == "중간 수집 (100개)":
                max_ch = 100
                st.info("⚡ 약 105 쿼터 소모 (약 8~10초 소요)")
            elif range_option == "직접 개수 입력":
                max_ch = st.number_input("수집할 최대 채널 수:", min_value=1, max_value=5000, value=25)
            else:
                max_ch = None
                st.info("💡 전체 구독 목록 전수 수집 및 분석")

        # Dynamic Button Label
        if not selected_channels:
            btn_label = "⚠️ 동기화할 채널을 먼저 선택하세요"
        elif len(selected_channels) == 1:
            ch_name = channel_map.get(selected_channels[0], "채널")
            btn_label = f"▶ '{ch_name}' 구독 목록 동기화 실행"
        else:
            btn_label = f"▶ 선택된 {len(selected_channels)}개 채널 통합 동기화 실행"

        if max_ch:
            btn_label += f" ({max_ch}개 부분 수집)"

        if st.button(
            btn_label,
            type="primary",
            disabled=(len(selected_channels) == 0),
            use_container_width=True,
        ):
            progress_bar = st.progress(0.0)
            status_placeholder = st.empty()

            try:
                total_synced = collector.run_sync_pipeline(
                    client_id=client_id,
                    client_secret=client_secret,
                    progress_callback=lambda p, msg: (
                        progress_bar.progress(min(1.0, max(0.0, p))),
                        status_placeholder.info(f"⏳ {msg}"),
                    ),
                    max_channels=max_ch,
                    target_channel_ids=selected_channels,
                )
                progress_bar.progress(1.0)
                status_placeholder.success(f"🎉 성공적으로 동기화가 완료되었습니다! (총 {total_synced}개 채널 수집 및 판정)")
                st.toast("동기화가 완료되었습니다. 사이드바와 대시보드가 업데이트됩니다.")
                st.rerun()
            except Exception as e:
                status_placeholder.error(f"❌ 동기화 중 오류가 발생했습니다: {str(e)}")

        # Summary Overview of Subscriptions in Tab 2
        st.markdown("---")
        st.markdown("#### 📊 현재 구독 데이터베이스 분류 현황")
        if not master_df.empty:
            sc1, sc2, sc3, sc4 = st.columns(4)
            sc1.metric("전체 수집 채널", f"{len(master_df):,}개")
            sc2.metric("📥 미분류 채널 (검토 대기)", f"{len(inbox_df):,}개")
            sc3.metric("💚 유지 채널 (확정)", f"{len(keep_df):,}개")
            sc4.metric("🗄️ 보류 보관함 (정리 대상)", f"{len(archived_df):,}개")

            if len(inbox_df) > 0:
                st.info(f"💡 현재 **{len(inbox_df):,}개**의 채널이 검토 대기 중입니다. 상단 **'3. 📥 미분류 채널 (Inbox)'** 탭에서 채널들을 손쉽게 분류해 보세요.")
            else:
                st.success("🎉 모든 채널의 분류가 완료되었습니다 (Inbox Zero 달성)!")
        else:
            st.info("ℹ️ 아직 수집된 채널이 없습니다. 위 동기화 버튼을 눌러 채널을 가져오세요.")


# ==========================================
# TAB 3: 미분류 채널 대기열 (Inbox Zero)
# ==========================================
with tab_inbox:
    st.subheader(f"📥 미분류 채널 대기열 (Inbox: {len(inbox_df):,}개 대기 중)")
    st.caption(
        "YouTube에서 새로 동기화된 채널들이 모이는 공간입니다. "
        "각 채널의 활동성을 확인하고 **[💚 유지]** 또는 **[📦 보류]**로 분류하여 대기열을 정리해 보세요 (Inbox Zero)."
    )

    if inbox_df.empty:
        st.success(
            "🎉 **미분류 채널이 없습니다 (Inbox Zero 달성)!**\n\n"
            "모든 구독 채널이 '유지' 또는 '보류'로 성공적으로 분류되었습니다.\n"
            "신규 계정을 연동하거나 '2. 수집 및 동기화'에서 새 채널을 가져오면 이곳에 자동으로 추가됩니다."
        )
    else:
        # Metrics KPI
        total_inbox = len(inbox_df)
        count_green = len(inbox_df[inbox_df["liveness_status"] == "GREEN"])
        count_yellow = len(inbox_df[inbox_df["liveness_status"] == "YELLOW"])
        count_red = len(inbox_df[inbox_df["liveness_status"] == "RED"])
        count_unknown = len(inbox_df[inbox_df["liveness_status"] == "UNKNOWN"])

        m1, m2, m3, m4, m5 = st.columns(5)
        m1.metric("검토 대기 채널", f"{total_inbox:,} 개")
        m2.metric("🟢 활성 (≤90일)", f"{count_green:,} 개")
        m3.metric("🟡 정체 (91~180일)", f"{count_yellow:,} 개")
        m4.metric("🔴 휴면 (>180일)", f"{count_red:,} 개")
        m5.metric("⚪ 미확인 (영상 0개)", f"{count_unknown:,} 개")

        st.divider()

        # Multi-filters
        col_f1, col_f2, col_f3, col_f4 = st.columns([2, 3, 2, 2])
        status_opts = ["GREEN", "YELLOW", "RED", "UNKNOWN"]

        selected_statuses = col_f1.multiselect(
            "신호등 상태 필터",
            options=status_opts,
            default=status_opts,
            format_func=lambda x: status_badge_map.get(x, x),
            key="inbox_status_filter",
        )
        search_keyword = col_f2.text_input("🔍 채널명 검색", placeholder="채널명을 입력하세요...", key="inbox_search_keyword")
        selected_category = col_f3.selectbox(
            "🏷️ 카테고리 필터",
            options=["전체"] + sorted(list(all_categories)),
            key="inbox_cat_filter",
        )
        selected_source = col_f4.selectbox(
            "👤 구독 소속 계정",
            options=["전체"] + sorted(list(all_sources)),
            key="inbox_src_filter",
        )

        # Apply filtering
        filtered_inbox = inbox_df[inbox_df["liveness_status"].isin(selected_statuses)]
        if search_keyword.strip():
            filtered_inbox = filtered_inbox[filtered_inbox["title"].str.contains(search_keyword.strip(), case=False, na=False)]
        if selected_category != "전체":
            filtered_inbox = filtered_inbox[filtered_inbox["categories"].apply(lambda cats: selected_category in cats)]
        if selected_source != "전체":
            filtered_inbox = filtered_inbox[filtered_inbox["source_accounts"].apply(lambda srcs: selected_source in srcs)]

        # Smart Batch Action
        with st.expander("⚡ 미분류 채널 스마트 일괄 관리 (원클릭 정리)"):
            col_b1, col_b2, col_b3 = st.columns(3)
            with col_b1:
                green_ids = filtered_inbox[filtered_inbox["liveness_status"] == "GREEN"]["channel_id"].tolist()
                if st.button(
                    f"💚 필터 활성(🟢) {len(green_ids)}개 일괄 유지",
                    disabled=(len(green_ids) == 0),
                    help="현재 검색/필터 결과 중 90일 이내에 업로드된 활성 채널들을 [유지]로 한 번에 이동합니다.",
                    key="batch_inbox_keep_green",
                    use_container_width=True,
                ):
                    db.set_channels_keep_batch(green_ids)
                    st.toast(f"{len(green_ids)}개 활성 채널이 유지 목록으로 이동되었습니다!", icon="💚")
                    st.rerun()
            with col_b2:
                red_ids = filtered_inbox[filtered_inbox["liveness_status"] == "RED"]["channel_id"].tolist()
                if st.button(
                    f"📦 필터 휴면(🔴) {len(red_ids)}개 일괄 보류",
                    disabled=(len(red_ids) == 0),
                    help="현재 검색/필터 결과 중 180일 이상 미업로드된 휴면 채널들을 [보류]로 한 번에 이동합니다.",
                    key="batch_inbox_arch_red",
                    use_container_width=True,
                ):
                    db.set_channels_archived_batch(red_ids, True)
                    st.toast(f"{len(red_ids)}개 휴면 채널이 보류 보관함으로 이동되었습니다!", icon="📦")
                    st.rerun()
            with col_b3:
                all_f_ids = filtered_inbox["channel_id"].tolist()
                if st.button(
                    f"💚 필터 채널 전체 ({len(all_f_ids)}개) 일괄 유지",
                    disabled=(len(all_f_ids) == 0),
                    help="현재 검색/필터된 모든 미분류 채널을 [유지]로 한 번에 이동합니다.",
                    key="batch_inbox_keep_all",
                    use_container_width=True,
                ):
                    db.set_channels_keep_batch(all_f_ids)
                    st.toast(f"{len(all_f_ids)}개 채널이 유지 목록으로 이동되었습니다!", icon="💚")
                    st.rerun()

        # Reset infinite scroll count if filters change
        current_inbox_hash = f"{selected_statuses}_{search_keyword}_{selected_category}_{selected_source}"
        if st.session_state.get("inbox_filter_hash") != current_inbox_hash:
            st.session_state["inbox_filter_hash"] = current_inbox_hash
            st.session_state["inbox_visible_count"] = 24

        view_col1, view_col2 = st.columns([2, 2])
        with view_col1:
            st.caption(f"검색/필터 결과: **{len(filtered_inbox):,}** 개 채널")
        with view_col2:
            view_mode = st.radio(
                "보기 방식:",
                ["🎴 카드 뷰 (Card Grid)", "📋 테이블 뷰 (Table)"],
                horizontal=True,
                label_visibility="collapsed",
                key="inbox_view_mode",
            )

        if view_mode == "🎴 카드 뷰 (Card Grid)":
            render_channel_cards_grid(filtered_inbox.to_dict(orient="records"), view_type="inbox", page_key_prefix="inbox")
        else:
            display_cols = [
                "liveness_status",
                "title",
                "subscriber_count_str",
                "video_count_str",
                "days_since_last_upload",
                "last_upload_at",
                "categories_str",
                "source_accounts_str",
                "custom_url",
            ]
            col_name_map = {
                "liveness_status": "상태",
                "title": "채널명",
                "subscriber_count_str": "구독자수",
                "video_count_str": "전체영상수",
                "days_since_last_upload": "미업로드 경과일수",
                "last_upload_at": "최근 업로드일",
                "categories_str": "카테고리",
                "source_accounts_str": "구독 소속 계정",
                "custom_url": "핸들/URL",
            }
            grid_df = filtered_inbox[display_cols].copy()
            grid_df["liveness_status"] = grid_df["liveness_status"].map(status_badge_map)
            grid_df.rename(columns=col_name_map, inplace=True)
            st.dataframe(grid_df, use_container_width=True, hide_index=True)

            # Table mode inspector dropdown
            st.divider()
            st.subheader("🔍 채널 상세 검사")
            channel_choices = filtered_inbox["title"].tolist()
            selected_ch_title = st.selectbox("상세 정보를 확인할 채널을 선택하세요:", channel_choices, key="tbl_sel_ch_inbox")
            if selected_ch_title:
                sel_row = filtered_inbox[filtered_inbox["title"] == selected_ch_title].iloc[0].to_dict()
                if st.button("📺 선택 채널 상세 팝업 열기", type="primary", key="btn_tbl_modal_inbox"):
                    show_channel_modal(sel_row)


# ==========================================
# TAB 4: 유지 채널 관리 (Keep List)
# ==========================================
with tab_keep:
    st.subheader(f"💚 유지 채널 관리 (Keep: {len(keep_df):,}개 보관 중)")
    st.caption("계속해서 시청하고 최신 영상을 챙겨볼 구독 확정 채널들입니다. 언제든 [↩️ 미분류] 또는 [📦 보류]로 상태를 변경할 수 있습니다.")

    if keep_df.empty:
        st.info("ℹ️ 현재 유지 목록으로 분류된 채널이 없습니다. 상단 **'3. 📥 미분류'** 탭에서 구독을 계속 유지할 채널의 **[💚 유지]** 버튼을 눌러보세요.")
    else:
        # Metrics KPI
        total_keep = len(keep_df)
        pct_keep = (total_keep / max(1, len(master_df))) * 100
        count_green = len(keep_df[keep_df["liveness_status"] == "GREEN"])
        count_yellow = len(keep_df[keep_df["liveness_status"] == "YELLOW"])
        count_red = len(keep_df[keep_df["liveness_status"] == "RED"])

        km1, km2, km3, km4, km5 = st.columns(5)
        km1.metric("확정 유지 채널", f"{total_keep:,} 개")
        km2.metric("전체 구독 대비 유지율", f"{pct_keep:.1f} %")
        km3.metric("🟢 활성 (≤90일)", f"{count_green:,} 개")
        km4.metric("🟡 정체 (91~180일)", f"{count_yellow:,} 개")
        km5.metric("🔴 휴면 (>180일)", f"{count_red:,} 개")

        st.divider()

        # Multi-filters
        col_kf1, col_kf2, col_kf3, col_kf4 = st.columns([2, 3, 2, 2])
        status_opts = ["GREEN", "YELLOW", "RED", "UNKNOWN"]

        k_selected_statuses = col_kf1.multiselect(
            "신호등 상태 필터",
            options=status_opts,
            default=status_opts,
            format_func=lambda x: status_badge_map.get(x, x),
            key="keep_status_filter",
        )
        k_search_keyword = col_kf2.text_input("🔍 채널명 검색", placeholder="채널명을 입력하세요...", key="keep_search_keyword")
        k_selected_category = col_kf3.selectbox(
            "🏷️ 카테고리 필터",
            options=["전체"] + sorted(list(all_categories)),
            key="keep_cat_filter",
        )
        k_selected_source = col_kf4.selectbox(
            "👤 구독 소속 계정",
            options=["전체"] + sorted(list(all_sources)),
            key="keep_src_filter",
        )

        filtered_keep = keep_df[keep_df["liveness_status"].isin(k_selected_statuses)]
        if k_search_keyword.strip():
            filtered_keep = filtered_keep[filtered_keep["title"].str.contains(k_search_keyword.strip(), case=False, na=False)]
        if k_selected_category != "전체":
            filtered_keep = filtered_keep[filtered_keep["categories"].apply(lambda cats: k_selected_category in cats)]
        if k_selected_source != "전체":
            filtered_keep = filtered_keep[filtered_keep["source_accounts"].apply(lambda srcs: k_selected_source in srcs)]

        # Smart Batch Action
        with st.expander("⚡ 유지 채널 스마트 일괄 관리"):
            col_kb1, col_kb2 = st.columns(2)
            with col_kb1:
                k_red_ids = filtered_keep[filtered_keep["liveness_status"] == "RED"]["channel_id"].tolist()
                if st.button(
                    f"📦 필터 휴면(🔴) {len(k_red_ids)}개 일괄 보류 이동",
                    disabled=(len(k_red_ids) == 0),
                    help="유지 목록 중 180일 이상 미업로드된 휴면 채널들을 보류 보관함으로 이동합니다.",
                    key="batch_keep_arch_red",
                    use_container_width=True,
                ):
                    db.set_channels_archived_batch(k_red_ids, True)
                    st.toast(f"{len(k_red_ids)}개 채널이 보류 보관함으로 이동되었습니다!", icon="📦")
                    st.rerun()
            with col_kb2:
                k_all_ids = filtered_keep["channel_id"].tolist()
                if st.button(
                    f"↩️ 현재 필터 채널 전체 ({len(k_all_ids)}개) 미분류 복귀",
                    disabled=(len(k_all_ids) == 0),
                    help="현재 필터된 유지 채널들을 다시 미분류(검토 대기열)로 복귀시킵니다.",
                    key="batch_keep_reset_inbox",
                    use_container_width=True,
                ):
                    db.set_channels_status_batch(k_all_ids, "INBOX")
                    st.toast(f"{len(k_all_ids)}개 채널이 미분류 상태로 복귀되었습니다!", icon="↩️")
                    st.rerun()

        current_keep_hash = f"{k_selected_statuses}_{k_search_keyword}_{k_selected_category}_{k_selected_source}"
        if st.session_state.get("keep_filter_hash") != current_keep_hash:
            st.session_state["keep_filter_hash"] = current_keep_hash
            st.session_state["keep_visible_count"] = 24

        k_view_col1, k_view_col2 = st.columns([2, 2])
        with k_view_col1:
            st.caption(f"검색/필터 결과: **{len(filtered_keep):,}** 개 채널")
        with k_view_col2:
            k_view_mode = st.radio(
                "보기 방식:",
                ["🎴 카드 뷰 (Card Grid)", "📋 테이블 뷰 (Table)"],
                horizontal=True,
                label_visibility="collapsed",
                key="keep_view_mode",
            )

        if k_view_mode == "🎴 카드 뷰 (Card Grid)":
            render_channel_cards_grid(filtered_keep.to_dict(orient="records"), view_type="keep", page_key_prefix="keep")
        else:
            display_cols = [
                "liveness_status",
                "title",
                "subscriber_count_str",
                "video_count_str",
                "days_since_last_upload",
                "last_upload_at",
                "categories_str",
                "source_accounts_str",
                "custom_url",
            ]
            col_name_map = {
                "liveness_status": "상태",
                "title": "채널명",
                "subscriber_count_str": "구독자수",
                "video_count_str": "전체영상수",
                "days_since_last_upload": "미업로드 경과일수",
                "last_upload_at": "최근 업로드일",
                "categories_str": "카테고리",
                "source_accounts_str": "구독 소속 계정",
                "custom_url": "핸들/URL",
            }
            grid_df = filtered_keep[display_cols].copy()
            grid_df["liveness_status"] = grid_df["liveness_status"].map(status_badge_map)
            grid_df.rename(columns=col_name_map, inplace=True)
            st.dataframe(grid_df, use_container_width=True, hide_index=True)

            st.divider()
            st.subheader("🔍 채널 상세 검사")
            channel_choices = filtered_keep["title"].tolist()
            selected_ch_title = st.selectbox("상세 정보를 확인할 채널을 선택하세요:", channel_choices, key="tbl_sel_ch_keep")
            if selected_ch_title:
                sel_row = filtered_keep[filtered_keep["title"] == selected_ch_title].iloc[0].to_dict()
                if st.button("📺 선택 채널 상세 팝업 열기", type="primary", key="btn_tbl_modal_keep"):
                    show_channel_modal(sel_row)


# ==========================================
# TAB 5: 보류 채널 보관함 (Archive Box)
# ==========================================
with tab_archive:
    st.subheader(f"🗄️ 보류 채널 보관함 (Archive: {len(archived_df):,}개 보관 중)")
    st.caption("더 이상 자주 보지 않거나 정리를 고려 중인 채널들을 임시 보관하는 공간입니다. 원할 때 언제든 **[💚 유지]** 또는 **[↩️ 미분류]**로 되돌릴 수 있습니다.")

    if archived_df.empty:
        st.info("ℹ️ 현재 보류 보관함에 보관된 채널이 없습니다. 정리가 필요한 채널 카드의 **[📦 보류]** 버튼을 눌러보세요.")
    else:
        # Archive Metrics
        am1, am2, am3 = st.columns(3)
        am1.metric("보류 채널 수", f"{len(archived_df):,} 개")
        am2.metric("전체 구독 대비 보류율", f"{(len(archived_df) / max(1, len(master_df)) * 100):.1f} %")
        am3.metric("유지 확정 채널 수", f"{len(keep_df):,} 개")

        with st.expander("⚡ 보관함 일괄 작업"):
            col_ab1, col_ab2 = st.columns(2)
            with col_ab1:
                if st.button("💚 보류 채널 전체 일괄 유지로 이동", type="secondary", key="batch_arch_to_keep_btn", use_container_width=True):
                    db.set_channels_keep_batch(archived_df["channel_id"].tolist())
                    st.toast(f"{len(archived_df)}개 보류 채널이 모두 유지 목록으로 이동되었습니다!", icon="💚")
                    st.rerun()
            with col_ab2:
                if st.button("↩️ 보류 채널 전체 일괄 미분류 복원", type="secondary", key="batch_restore_all_btn", use_container_width=True):
                    db.set_channels_status_batch(archived_df["channel_id"].tolist(), "INBOX")
                    st.toast(f"{len(archived_df)}개 보류 채널이 모두 미분류 상태로 복원되었습니다!", icon="♻️")
                    st.rerun()

        # Filter controls for archive
        col_af1, col_af2, col_af3, col_af4 = st.columns([2, 3, 2, 2])
        status_opts = ["GREEN", "YELLOW", "RED", "UNKNOWN"]

        arch_statuses = col_af1.multiselect(
            "신호등 상태 필터",
            options=status_opts,
            default=status_opts,
            format_func=lambda x: status_badge_map.get(x, x),
            key="arch_status_filter",
        )
        arch_keyword = col_af2.text_input("🔍 보류 채널명 검색", placeholder="채널명을 입력하세요...", key="arch_search_keyword")
        arch_category = col_af3.selectbox(
            "🏷️ 카테고리 필터",
            options=["전체"] + sorted(list(all_categories)),
            key="arch_cat_filter",
        )
        arch_source = col_af4.selectbox(
            "👤 구독 소속 계정",
            options=["전체"] + sorted(list(all_sources)),
            key="arch_src_filter",
        )

        arch_filtered = archived_df[archived_df["liveness_status"].isin(arch_statuses)]
        if arch_keyword.strip():
            arch_filtered = arch_filtered[arch_filtered["title"].str.contains(arch_keyword.strip(), case=False, na=False)]
        if arch_category != "전체":
            arch_filtered = arch_filtered[arch_filtered["categories"].apply(lambda cats: arch_category in cats)]
        if arch_source != "전체":
            arch_filtered = arch_filtered[arch_filtered["source_accounts"].apply(lambda srcs: arch_source in srcs)]

        current_arch_hash = f"{arch_statuses}_{arch_keyword}_{arch_category}_{arch_source}"
        if st.session_state.get("arch_filter_hash") != current_arch_hash:
            st.session_state["arch_filter_hash"] = current_arch_hash
            st.session_state["archive_visible_count"] = 24

        aview_col1, aview_col2 = st.columns([2, 2])
        with aview_col1:
            st.caption(f"보류 보관함 결과: **{len(arch_filtered):,}** 개 채널")
        with aview_col2:
            arch_view_mode = st.radio(
                "보기 방식:",
                ["🎴 카드 뷰 (Card Grid)", "📋 테이블 뷰 (Table)"],
                horizontal=True,
                label_visibility="collapsed",
                key="arch_view_mode",
            )

        if arch_view_mode == "🎴 카드 뷰 (Card Grid)":
            render_channel_cards_grid(arch_filtered.to_dict(orient="records"), view_type="archive", page_key_prefix="archive")
        else:
            display_cols = [
                "liveness_status",
                "title",
                "subscriber_count_str",
                "video_count_str",
                "days_since_last_upload",
                "last_upload_at",
                "categories_str",
                "source_accounts_str",
                "custom_url",
            ]
            col_name_map = {
                "liveness_status": "상태",
                "title": "채널명",
                "subscriber_count_str": "구독자수",
                "video_count_str": "전체영상수",
                "days_since_last_upload": "미업로드 경과일수",
                "last_upload_at": "최근 업로드일",
                "categories_str": "카테고리",
                "source_accounts_str": "구독 소속 계정",
                "custom_url": "핸들/URL",
            }
            agrid_df = arch_filtered[display_cols].copy()
            agrid_df["liveness_status"] = agrid_df["liveness_status"].map(status_badge_map)
            agrid_df.rename(columns=col_name_map, inplace=True)
            st.dataframe(agrid_df, use_container_width=True, hide_index=True)

            st.divider()
            st.subheader("🔍 보류 채널 상세 검사")
            arch_choices = arch_filtered["title"].tolist()
            selected_arch_title = st.selectbox("상세 정보를 확인할 채널을 선택하세요:", arch_choices, key="tbl_sel_ch_arch")
            if selected_arch_title:
                sel_arch_row = arch_filtered[arch_filtered["title"] == selected_arch_title].iloc[0].to_dict()
                if st.button("📺 선택 채널 상세 팝업 열기", type="primary", key="btn_tbl_modal_arch"):
                    show_channel_modal(sel_arch_row)


# ==========================================
# TAB 6: 엑셀 내보내기 (Export Engine)
# ==========================================
with tab_export:
    st.subheader("📥 엑셀(.xlsx) 리포트 추출")
    st.write("통합 구독 목록을 상태별 조건부 서식과 상세 정보가 정돈된 엑셀 파일로 내려받아 영구 보존할 수 있습니다.")

    if master_df.empty:
        st.info("내보낼 구독 채널 데이터가 없습니다.")
    else:
        export_scope = st.radio(
            "내보낼 구독 상태 범위:",
            [
                "전체 채널 (미분류 + 유지 + 보류)",
                "📥 미분류 채널만 (Inbox)",
                "💚 유지 채널만 (Keep)",
                "🗄️ 보류 보관함 채널만 (Archive)",
            ],
            horizontal=True,
        )

        if export_scope == "📥 미분류 채널만 (Inbox)":
            base_records = [r for r in records if r["review_status"] == "INBOX"]
        elif export_scope == "💚 유지 채널만 (Keep)":
            base_records = [r for r in records if r["review_status"] == "KEEP"]
        elif export_scope == "🗄️ 보류 보관함 채널만 (Archive)":
            base_records = [r for r in records if r["review_status"] == "ARCHIVED"]
        else:
            base_records = records

        target_records = base_records
        st.write(f"추출 대상: **{len(target_records):,}** 개 채널")

        if st.button("📊 엑셀 파일 생성 및 준비", type="primary"):
            rows = []
            status_trans_map = {"INBOX": "미분류", "KEEP": "유지", "ARCHIVED": "보류"}
            for r in target_records:
                vids = json.loads(r["recent_videos"]) if r["recent_videos"] else []
                sources = json.loads(r["source_accounts"]) if r["source_accounts"] else []
                cats = json.loads(r["categories"]) if r["categories"] else []
                raw_ch = json.loads(r["raw_channel_json"]) if r["raw_channel_json"] else {}
                stats = raw_ch.get("statistics", {})
                c_url = f"https://www.youtube.com/{r['custom_url']}" if r["custom_url"] else f"https://www.youtube.com/channel/{r['channel_id']}"

                sub_cnt = stats.get("subscriberCount")
                vid_cnt = stats.get("videoCount")
                rev = r["review_status"] or "INBOX"

                rows.append({
                    "분류상태": status_trans_map.get(rev, rev),
                    "활동상태": r["liveness_status"],
                    "채널명": r["title"],
                    "구독자수": int(sub_cnt) if sub_cnt and str(sub_cnt).isdigit() else (sub_cnt or "-"),
                    "전체영상수": int(vid_cnt) if vid_cnt and str(vid_cnt).isdigit() else (vid_cnt or "-"),
                    "채널홈URL": c_url,
                    "카테고리": ", ".join(cats),
                    "최근 업로드일": r["last_upload_at"] or "-",
                    "미업로드 경과일수": r["days_since_last_upload"] if r["days_since_last_upload"] is not None else "-",
                    "구독계정출처": ", ".join(sources),
                    "최근영상1": vids[0]["title"] if len(vids) > 0 else "",
                    "최근영상2": vids[1]["title"] if len(vids) > 1 else "",
                    "최근영상3": vids[2]["title"] if len(vids) > 2 else "",
                    "최근영상4": vids[3]["title"] if len(vids) > 3 else "",
                    "최근영상5": vids[4]["title"] if len(vids) > 4 else "",
                    "채널설명": (r["description"] or "").replace("\n", " "),
                })

            export_df = pd.DataFrame(rows)

            # Build openpyxl workbook with styling
            output = io.BytesIO()
            with pd.ExcelWriter(output, engine="openpyxl") as writer:
                export_df.to_excel(writer, index=False, sheet_name="Subscriptions_SSOT")
                ws = writer.sheets["Subscriptions_SSOT"]

                # Header styles
                header_font = Font(name="Malgun Gothic", size=11, bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
                for cell in ws[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Conditional formatting for Status column (Column B: 활동상태)
                green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                green_font = Font(color="006100", bold=True)
                yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                yellow_font = Font(color="9C6500", bold=True)
                red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                red_font = Font(color="9C0006", bold=True)

                for row in range(2, len(rows) + 2):
                    status_cell = ws.cell(row=row, column=2)
                    val = str(status_cell.value)
                    if val == "GREEN":
                        status_cell.fill = green_fill
                        status_cell.font = green_font
                    elif val == "YELLOW":
                        status_cell.fill = yellow_fill
                        status_cell.font = yellow_font
                    elif val == "RED":
                        status_cell.fill = red_fill
                        status_cell.font = red_font
                    status_cell.alignment = Alignment(horizontal="center", vertical="center")

                # Auto-fit column widths
                for col in ws.columns:
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    max_len = max(len(str(cell.value or "")) for cell in col)
                    ws.column_dimensions[col_letter].width = min(50, max(max_len + 3, 12))

            excel_data = output.getvalue()
            filename = f"youtube_subscriptions_ssot_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

            st.download_button(
                label="📥 엑셀(.xlsx) 파일 다운로드",
                data=excel_data,
                file_name=filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
            )

